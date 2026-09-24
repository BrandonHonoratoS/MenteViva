"""Exportes a Excel con identidad Mente Viva (openpyxl): tablero organizacional / de área."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..agents import catalogo as C

VIOLETA = "6D28D9"


def _hoja(wb, titulo: str, encabezados: list[str], filas: list[list], primera: bool = False):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo[:31]
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VIOLETA)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for f in filas:
        ws.append(f)
    for i, _ in enumerate(encabezados, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(60, max((len(str(r[i - 1])) for r in filas), default=10) + 2))
    ws.freeze_panes = "A2"
    return ws


def tablero(k: dict, empresa: str, alcance: str) -> bytes:
    wb = Workbook()
    resumen = [["Mente Viva · Tablero", empresa, alcance], ["Periodo (días)", k["periodo_dias"], ""], ["Colaboradores", k["colaboradores"], ""], ["Activos", k["activos"], ""],
               ["Sesiones en el periodo", k["sesiones_periodo"], ""], ["", "", ""], ["Índice", "Valor", "Fórmula"]]
    for i in k["indices"].values():
        resumen.append([i["nombre"], f"{i['valor']}{i['unidad']}" if i["valor"] is not None else "sin datos", i["formula"]])
    _hoja(wb, "Resumen", ["Concepto", "Valor", "Detalle"], resumen, primera=True)
    _hoja(wb, "Colaboradores", ["Nombre", "Área", "Puesto", "Diagnóstico", "Ventas", "Entrevistas", "Ruta DM", "Sesiones", "Score prom.", "Tendencia", "Semáforo", "Última actividad", "Racha (sem)"],
          [[r["nombre"], r["area"], r["puesto"], r["diagnostico"], r["niveles"].get("ventas", ""), r["niveles"].get("entrevistas", ""), r["niveles"].get("ruta_dm", ""), r["sesiones"],
            r["score_promedio"], r["tendencia"], r["semaforo"], r["ultima_actividad"], r["racha"]] for r in k["colaboradores_resumen"]])
    if k.get("por_area"):
        _hoja(wb, "Áreas", ["Área", "Director", "Colaboradores", "Activos", "Sesiones", "Score", "Requieren atención"],
              [[a["nombre"], a["director"], a["colaboradores"], a["activos"], a["sesiones"], a["score"], a["atencion"]] for a in k["por_area"]])
    _hoja(wb, "Ruta DM", ["Nivel", "Colaboradores", "Valor del nivel"], [[n, c, C.valor_dm(n)] for n, c in k["distribucion_dm"].items()])
    _hoja(wb, "Serie semanal", ["Semana", "Sesiones", "Score promedio"], [[s["semana"], s["sesiones"], s["score"]] for s in k["serie"]])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sesiones(lista: list[dict]) -> bytes:
    wb = Workbook()
    _hoja(wb, "Sesiones", ["ID", "Fecha", "Colaborador", "Habilidad", "Nivel", "Competencia", "Score", "Turnos", "Objetivo", "Voluntaria"],
          [[s["id"], s["inicio"][:16].replace("T", " "), s["usuario_nombre"], C.HABILIDADES.get(s["habilidad"], {}).get("corto", s["habilidad"]), s.get("nivel"), s.get("competencia"),
            s.get("score_global"), s.get("turnos"), s.get("objetivo"), "sí" if s.get("voluntaria") else "no"] for s in lista], primera=True)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
