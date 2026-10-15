"""Encuesta de primera experiencia: se pide una sola vez tras la primera práctica; RRHH la ve sin precio, el superadmin con precio; Excel."""
from __future__ import annotations

import pytest

from app import db
from tests.conftest import ENCUESTA_VALIDA, Sesion
from tests.test_laboratorio import FEEDBACK_LAB, _colaborador_con_diagnostico


@pytest.mark.encuesta_obligatoria
def test_encuesta_tras_primera_practica(client, stub):
    rocio = _colaborador_con_diagnostico(client, stub, "rocio@condor.mx", "Rocío Encuesta")
    uid = rocio.u["id"]
    assert db.pendiente_encuesta(uid) is None and "Contar mi experiencia" not in rocio.get("/hoy").text
    # primera práctica (laboratorio) completada → aparece la invitación en Hoy y en el reporte
    r = rocio.post("/laboratorio/iniciar")
    sid = int(r.headers["location"].rsplit("/", 1)[1])
    rocio.api(f"/api/sesiones/{sid}/mensaje", {"texto": "Listo"})
    stub["juan.laboratorio"] = {"mensaje": "El cliente asiente.", "accion": "avanzar", "tension": 40}
    rocio.api(f"/api/sesiones/{sid}/mensaje", {"texto": "Pido evidencia a QA."})
    stub["juan.lab_feedback"] = FEEDBACK_LAB
    rocio.api(f"/api/sesiones/{sid}/mensaje", {"texto": "Fin"})
    assert db.sesion(sid)["estado"] == "completada"
    assert db.pendiente_encuesta(uid)["id"] == sid
    assert "Contar mi experiencia" in rocio.get("/hoy").text and "Contar mi experiencia" in rocio.get(f"/sesion/{sid}/reporte").text
    # obligatoria: no se puede abrir otra práctica hasta responderla
    r = rocio.get("/entrenar/ventas")
    assert r.status_code == 307 and r.headers["location"].startswith("/encuesta?obligatoria=1")
    assert rocio.get("/laboratorio").status_code == 307 and rocio.post("/entrenar/ventas/iniciar").status_code == 409
    r = rocio.get("/encuesta?obligatoria=1")
    assert r.status_code == 200 and "Valor y disposición de pago" in r.text and "Antes de tu siguiente práctica" in r.text and "¿Cuántas veces has utilizado Mente Viva?" in r.text
    # incompleta → se marcan faltantes; completa → se guarda una sola vez
    r = rocio.post("/encuesta", {"calificacion": "9"})
    assert r.status_code == 400 and "Faltan respuestas" in r.text
    datos = dict(ENCUESTA_VALIDA)
    r = rocio.post("/encuesta", datos)
    assert r.status_code == 303
    e = db.encuesta_de(uid)
    assert e and e["respuestas"]["calificacion"] == 9 and e["respuestas"]["precio_unico"] == 1500.0 and e["respuestas"]["funciones"] == ["Avatares con voz", "App móvil"] and e["sesion_id"] == sid
    assert "¡Gracias" in rocio.get("/encuesta").text and db.pendiente_encuesta(uid) is None
    assert "Contar mi experiencia" not in rocio.get("/hoy").text
    assert rocio.get("/entrenar/ventas").status_code == 200   # ya puede seguir practicando
    assert rocio.post("/encuesta", datos).status_code == 303 and len(db.encuestas(rocio.u["empresa_id"])) >= 1
    # RRHH ve experiencia sin precio; superadmin ve precio; Excel en ambos
    rrhh = Sesion(client, "rrhh@condor.mx", "NuevaClaveRRHH1")
    r = rrhh.get("/encuestas")
    assert r.status_code == 200 and "El feedback con evidencia" in r.text and "1,500" not in r.text and "$150 a $199 MXN" not in r.text
    eid = e["id"]
    r = rrhh.get(f"/encuestas/{eid}")
    assert r.status_code == 200 and "¿Cuántas veces has utilizado Mente Viva?" in r.text and "Retroalimentación de la IA" in r.text and "Valor y disposición de pago" not in r.text
    admin = Sesion(client, "admin@menteviva.mx", "SuperClave123")
    r = admin.get("/admin/encuestas")
    assert r.status_code == 200 and "$150 a $199 MXN" in r.text and "Nueva encuesta" in str([n["titulo"] for n in db.notificaciones(admin.u["id"])])
    r = admin.get(f"/admin/encuestas/{eid}")
    assert r.status_code == 200 and "Valor y disposición de pago" in r.text and "Mensual" in r.text and "De inmediato" in r.text
    for s, url in ((rrhh, "/exportar/encuestas.xlsx"), (admin, "/admin/encuestas.xlsx")):
        r = s.get(url)
        assert r.status_code == 200 and r.content[:2] == b"PK"
    from openpyxl import load_workbook
    from io import BytesIO
    ws = load_workbook(BytesIO(rrhh.get("/exportar/encuestas.xlsx").content)).active
    encabezados = [c.value for c in ws[1]]
    assert "¿Qué fue lo que más te gustó?" in encabezados and not any("dispuesto" in (h or "").lower() for h in encabezados)
    ws = load_workbook(BytesIO(admin.get("/admin/encuestas.xlsx").content)).active
    assert any("dispuesto" in (c.value or "").lower() for c in ws[1])
